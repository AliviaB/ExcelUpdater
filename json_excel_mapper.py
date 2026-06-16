"""
json_excel_mapper.py
--------------------
Reads attribute names from Column A of an Excel file, searches a JSON file
for each attribute key anywhere in its hierarchy, and writes:
  - Column B: the full dot-notation path  (e.g. data.team.memberId)
  - Column C: the extracted value

Usage:
    python json_excel_mapper.py                          # uses defaults
    python json_excel_mapper.py data.json mapping.xlsx  # custom paths
    python json_excel_mapper.py data.json mapping.xlsx Sheet2  # custom sheet
"""

import json
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ── helpers ──────────────────────────────────────────────────────────────────

def find_all_paths(obj, target_key, current_path=""):
    """
    Recursively walk the JSON object and collect every dot-notation path
    whose final key matches target_key (case-insensitive).
    Returns a list of (path_string, value) tuples.
    """
    results = []

    if isinstance(obj, dict):
        for k, v in obj.items():
            path = f"{current_path}.{k}" if current_path else k
            if k.lower() == target_key.lower():
                results.append((path, v))
            # keep descending even after a match (key may appear at multiple levels)
            results.extend(find_all_paths(v, target_key, path))

    elif isinstance(obj, list):
        for idx, item in enumerate(obj):
            path = f"{current_path}[{idx}]"
            results.extend(find_all_paths(item, target_key, path))

    return results


def format_value(value):
    """Flatten a value to a readable string for the cell."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value) if value is not None else ""


# ── styling helpers ───────────────────────────────────────────────────────────

def _apply_found(cell):
    cell.fill = PatternFill("solid", start_color="DFF2E1")   # light green
    cell.font = Font(name="Arial", size=10, color="1A5C2A")

def _apply_multi(cell):
    cell.fill = PatternFill("solid", start_color="FFF3CD")   # amber
    cell.font = Font(name="Arial", size=10, color="7A5000")

def _apply_missing(cell):
    cell.fill = PatternFill("solid", start_color="FAD7D7")   # light red
    cell.font = Font(name="Arial", size=10, color="8B0000")

def _plain(cell):
    cell.font = Font(name="Arial", size=10)


# ── main logic ────────────────────────────────────────────────────────────────

def process(json_path: str, excel_path: str, sheet_name: str | None = None):
    # Load JSON
    json_file = Path(json_path)
    if not json_file.exists():
        raise FileNotFoundError(f"JSON file not found: {json_path}")
    with open(json_file, encoding="utf-8") as f:
        json_data = json.load(f)

    # Load Excel
    excel_file = Path(excel_path)
    if not excel_file.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")
    wb = load_workbook(excel_file)
    ws = wb[sheet_name] if sheet_name else wb.active

    # Ensure header row labels (B1, C1) exist
    if ws["B1"].value is None:
        ws["B1"] = "JSON Key Path"
        ws["B1"].font = Font(bold=True, name="Arial")
    if ws["C1"].value is None:
        ws["C1"] = "Extracted Value"
        ws["C1"].font = Font(bold=True, name="Arial")

    found = missing = multi = 0

    # Iterate attribute names in column A (skip header row 1)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        attr_cell = row[0]          # Column A
        path_cell = row[1]          # Column B
        value_cell = row[2]         # Column C

        attr_name = attr_cell.value
        if not attr_name:
            continue
        attr_name = str(attr_name).strip()

        matches = find_all_paths(json_data, attr_name)

        if not matches:
            path_cell.value = "NOT FOUND"
            value_cell.value = ""
            _apply_missing(path_cell)
            _apply_missing(value_cell)
            missing += 1

        elif len(matches) == 1:
            path, val = matches[0]
            path_cell.value = path
            value_cell.value = format_value(val)
            _apply_found(path_cell)
            _apply_found(value_cell)
            found += 1

        else:
            # Multiple occurrences – write all paths separated by " | "
            paths  = " | ".join(m[0] for m in matches)
            values = " | ".join(format_value(m[1]) for m in matches)
            path_cell.value  = paths
            value_cell.value = values
            _apply_multi(path_cell)
            _apply_multi(value_cell)
            multi += 1

        # Wrap long text
        for cell in (path_cell, value_cell):
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Auto-fit columns B and C (rough heuristic)
    for col_letter, col_idx in (("B", 2), ("C", 3)):
        max_len = max(
            (len(str(ws.cell(r, col_idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    output_path = excel_file.parent / f"mapped_{excel_file.name}"
    wb.save(output_path)

    print(f"\n{'─'*50}")
    print(f"  ✅  Matched  : {found}")
    print(f"  ⚠️   Multiple : {multi}")
    print(f"  ❌  Missing  : {missing}")
    print(f"  📄  Output   : {output_path}")
    print(f"{'─'*50}\n")
    return str(output_path)


# ── entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    args = sys.argv[1:]
    json_path  = args[0] if len(args) > 0 else "sample_data.json"
    excel_path = args[1] if len(args) > 1 else "mapping_template.xlsx"
    sheet      = args[2] if len(args) > 2 else None

    process(json_path, excel_path, sheet)
